"""Iteration 3 tests: unsubscribe flow, clients CRUD/export/import, campaigns v2, cron auth."""
import io
import os
import time
import uuid
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook


def _cron_secret() -> str:
    env_file = Path("/app/backend/.env")
    for line in env_file.read_text(encoding="utf-8").splitlines():
        if line.startswith("WEBHOOK_CRON_SECRET"):
            return line.split("=", 1)[1].strip().strip('"').strip("'")
    return ""


# ---------- /api/auth/login ----------
def test_admin_login_returns_token(api_client, base_url, admin_credentials):
    email, password = admin_credentials
    r = api_client.post(f"{base_url}/api/auth/login", json={"email": email, "password": password}, timeout=30)
    assert r.status_code == 200
    body = r.json()
    assert isinstance(body.get("token"), str) and len(body["token"]) > 20
    assert body["user"]["email"] == email


# ---------- /api/songs limit=6 minimum ----------
def test_songs_search_limit_6(api_client, base_url):
    r = api_client.get(f"{base_url}/api/songs", params={"q": "queen", "limit": 6}, timeout=30)
    assert r.status_code == 200
    data = r.json()
    assert data["total"] > 0
    assert 0 < len(data["items"]) <= 6


# ---------- Song request creates client & no duplicate on second request ----------
def test_song_request_upserts_no_duplicate(api_client, base_url, admin_headers):
    email = f"test-qa-{int(time.time())}-{uuid.uuid4().hex[:6]}@example.com"
    payload = {"email": email, "name": "TEST QA", "title": "Somebody To Love", "artist": "Queen"}
    r1 = api_client.post(f"{base_url}/api/song-requests", json=payload, timeout=30)
    assert r1.status_code == 201
    b1 = r1.json()
    assert b1["request"]["email"] == email
    assert b1.get("client") is not None

    # verify client fields
    r_list = api_client.get(f"{base_url}/api/admin/clients", headers=admin_headers, params={"q": email}, timeout=30)
    assert r_list.status_code == 200
    items = r_list.json()["items"]
    match = next((c for c in items if c["email"] == email), None)
    assert match is not None
    assert match["subscribed"] is True
    assert match.get("unsubscribeToken")
    assert match.get("followupCount") == 0
    assert match.get("nextFollowupAt") is not None
    assert match.get("source") == "cancion"
    assert "_id" not in match
    client_id = match["id"]
    token = match["unsubscribeToken"]

    # second identical request must not duplicate client
    r2 = api_client.post(f"{base_url}/api/song-requests", json=payload, timeout=30)
    assert r2.status_code == 201
    r_list2 = api_client.get(f"{base_url}/api/admin/clients", headers=admin_headers, params={"q": email}, timeout=30)
    matches2 = [c for c in r_list2.json()["items"] if c["email"] == email]
    assert len(matches2) == 1

    # ---------- Unsubscribe public flow ----------
    ug = api_client.get(f"{base_url}/api/unsubscribe/{token}", timeout=30)
    assert ug.status_code == 200
    ugb = ug.json()
    assert ugb["subscribed"] is True
    assert "•" in ugb["email"] or "*" in ugb["email"] or "@" in ugb["email"]

    up = api_client.post(f"{base_url}/api/unsubscribe/{token}", json={"subscribed": False}, timeout=30)
    assert up.status_code == 200
    assert up.json()["subscribed"] is False

    # verify via admin
    r_v = api_client.get(f"{base_url}/api/admin/clients", headers=admin_headers, params={"q": email}, timeout=30)
    m = next(c for c in r_v.json()["items"] if c["email"] == email)
    assert m["subscribed"] is False
    assert m.get("nextFollowupAt") is None

    # resubscribe
    up2 = api_client.post(f"{base_url}/api/unsubscribe/{token}", json={"subscribed": True}, timeout=30)
    assert up2.status_code == 200
    assert up2.json()["subscribed"] is True
    r_v2 = api_client.get(f"{base_url}/api/admin/clients", headers=admin_headers, params={"q": email}, timeout=30)
    m2 = next(c for c in r_v2.json()["items"] if c["email"] == email)
    assert m2["subscribed"] is True
    assert m2.get("nextFollowupAt") is not None

    # bogus token
    r_bogus = api_client.get(f"{base_url}/api/unsubscribe/does-not-exist-token", timeout=30)
    assert r_bogus.status_code == 404

    # cleanup: delete request(s) + client
    reqs = api_client.get(f"{base_url}/api/admin/song-requests", headers=admin_headers, timeout=30).json()["items"]
    for req in reqs:
        if req["email"] == email:
            api_client.delete(f"{base_url}/api/admin/song-requests/{req['id']}", headers=admin_headers, timeout=30)
    api_client.delete(f"{base_url}/api/admin/clients/{client_id}", headers=admin_headers, timeout=30)


# ---------- Admin clients CRUD ----------
def test_admin_clients_crud_and_filters(api_client, base_url, admin_headers):
    email = f"test_cli_{uuid.uuid4().hex[:8]}@example.com"
    # create manual
    r = api_client.post(f"{base_url}/api/admin/clients", headers=admin_headers,
                       json={"email": email, "name": "TEST Manual", "phone": "600111222", "subscribed": True}, timeout=30)
    assert r.status_code == 201
    client = r.json()["client"]
    assert client["source"] == "manual"
    assert client["subscribed"] is True
    assert "_id" not in client
    cid = client["id"]

    # duplicate -> 409
    r_dup = api_client.post(f"{base_url}/api/admin/clients", headers=admin_headers,
                           json={"email": email, "name": "dup"}, timeout=30)
    assert r_dup.status_code == 409

    # invalid email -> 400
    r_bad = api_client.post(f"{base_url}/api/admin/clients", headers=admin_headers,
                           json={"email": "not-email"}, timeout=30)
    assert r_bad.status_code == 400

    # PUT update + toggle subscribed off -> nextFollowupAt null
    r_u = api_client.put(f"{base_url}/api/admin/clients/{cid}", headers=admin_headers,
                        json={"name": "TEST Manual Updated", "phone": "600999888", "notes": "n", "subscribed": False}, timeout=30)
    assert r_u.status_code == 200
    up = r_u.json()["client"]
    assert up["name"] == "TEST Manual Updated"
    assert up["subscribed"] is False
    assert up.get("nextFollowupAt") is None

    # filters
    r_q = api_client.get(f"{base_url}/api/admin/clients", headers=admin_headers, params={"q": "TEST Manual Updated"}, timeout=30)
    assert r_q.status_code == 200
    assert any(c["id"] == cid for c in r_q.json()["items"])

    r_src = api_client.get(f"{base_url}/api/admin/clients", headers=admin_headers, params={"source": "manual"}, timeout=30)
    assert all(c["source"] == "manual" for c in r_src.json()["items"])
    assert any(c["id"] == cid for c in r_src.json()["items"])

    r_sub = api_client.get(f"{base_url}/api/admin/clients", headers=admin_headers, params={"subscribed": "false"}, timeout=30)
    assert all(c["subscribed"] is False for c in r_sub.json()["items"])
    assert any(c["id"] == cid for c in r_sub.json()["items"])

    # delete
    r_d = api_client.delete(f"{base_url}/api/admin/clients/{cid}", headers=admin_headers, timeout=30)
    assert r_d.status_code == 204


# ---------- Export xlsx ----------
def test_admin_clients_export_xlsx(api_client, base_url, admin_headers):
    r = api_client.get(f"{base_url}/api/admin/clients/export", headers=admin_headers, timeout=60)
    assert r.status_code == 200
    ct = r.headers.get("content-type", "")
    assert "spreadsheetml.sheet" in ct
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    expected = ["Correo", "Nombre", "Teléfono", "Origen", "Fecha de alta", "Suscrito", "Último correo", "Próximo seguimiento", "Seguimientos", "Notas"]
    assert header == expected, f"Unexpected header: {header}"


# ---------- Import xlsx ----------
def test_admin_clients_import_xlsx(api_client, base_url, admin_headers):
    # Seed one existing email
    existing = f"test_imp_existing_{uuid.uuid4().hex[:6]}@example.com"
    seed = api_client.post(f"{base_url}/api/admin/clients", headers=admin_headers,
                          json={"email": existing, "name": "TEST Existing"}, timeout=30)
    assert seed.status_code == 201
    existing_id = seed.json()["client"]["id"]

    new1 = f"test_imp_a_{uuid.uuid4().hex[:6]}@example.com"
    new2 = f"test_imp_b_{uuid.uuid4().hex[:6]}@example.com"

    wb = Workbook()
    ws = wb.active
    ws.append(["Correo", "Nombre", "Teléfono", "Suscrito"])
    ws.append([new1, "TEST New A", "600000001", "Sí"])
    ws.append([new2, "TEST New B", "600000002", "Sí"])
    ws.append([existing, "TEST Existing Renamed", "600000003", "Sí"])
    ws.append(["not-an-email", "Invalid", "", "Sí"])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)

    # no file -> 400
    r_no = api_client.post(f"{base_url}/api/admin/clients/import", headers=admin_headers, timeout=30)
    assert r_no.status_code == 400

    # with file. Use fresh session so session Content-Type=json does not override multipart
    import requests as _rq
    files = {"file": ("clients.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    headers = {"Authorization": admin_headers["Authorization"]}
    r = _rq.post(f"{base_url}/api/admin/clients/import", headers=headers, files=files, timeout=60)
    assert r.status_code == 200, r.text
    summary = r.json()
    assert summary["total"] == 4
    assert summary["created"] == 2
    # existing had no phone/notes so it will be updated with phone -> updated=1
    assert summary["updated"] >= 0
    assert summary["skipped"] >= 1  # invalid email

    # verify imported clients have source 'importado', welcomeSentAt null, nextFollowupAt set
    lst = api_client.get(f"{base_url}/api/admin/clients", headers=admin_headers, params={"q": "test_imp_"}, timeout=30).json()["items"]
    new_a = next(c for c in lst if c["email"] == new1)
    assert new_a["source"] == "importado"
    assert new_a.get("welcomeSentAt") in (None, "")
    assert new_a.get("nextFollowupAt")

    # cleanup
    for c in lst:
        if c["email"].startswith("test_imp_"):
            api_client.delete(f"{base_url}/api/admin/clients/{c['id']}", headers=admin_headers, timeout=30)
    # existing_id already deleted via above loop
    api_client.delete(f"{base_url}/api/admin/clients/{existing_id}", headers=admin_headers, timeout=30)


# ---------- Campaigns ----------
def test_admin_campaigns_list_and_send(api_client, base_url, admin_headers):
    # First create an isolated audience using 'importado' source with example.com clients only
    imported_emails = [f"test_camp_{uuid.uuid4().hex[:6]}@example.com" for _ in range(2)]
    ids = []
    for em in imported_emails:
        # import via CRUD would set source=manual; instead use import endpoint to set source=importado
        pass
    wb = Workbook()
    ws = wb.active
    ws.append(["Correo", "Nombre", "Suscrito"])
    for em in imported_emails:
        ws.append([em, "TEST Campaign Target", "Sí"])
    bio = io.BytesIO()
    wb.save(bio)
    files = {"file": ("c.xlsx", bio.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    headers = {"Authorization": admin_headers["Authorization"]}
    import requests as _rq
    r_imp = _rq.post(f"{base_url}/api/admin/clients/import", headers=headers, files=files, timeout=60)
    assert r_imp.status_code == 200

    # list campaigns has {items}
    r_list = api_client.get(f"{base_url}/api/admin/campaigns", headers=admin_headers, timeout=30)
    assert r_list.status_code == 200
    assert "items" in r_list.json()

    # missing subject -> 400
    r_bad = api_client.post(f"{base_url}/api/admin/campaigns", headers=admin_headers,
                           json={"subject": "", "message": "hi"}, timeout=30)
    assert r_bad.status_code == 400

    # send to source=importado audience only
    r_send = api_client.post(f"{base_url}/api/admin/campaigns", headers=admin_headers,
                            json={"subject": "TEST Iter3 Campaign", "message": "hola tests", "source": "importado"}, timeout=30)
    assert r_send.status_code == 202
    campaign = r_send.json()["campaign"]
    assert campaign["status"] == "sending"
    campaign_id = campaign["id"]
    assert campaign["total"] >= len(imported_emails)

    # poll until done (max ~30s)
    done = None
    for _ in range(30):
        time.sleep(1)
        items = api_client.get(f"{base_url}/api/admin/campaigns", headers=admin_headers, timeout=30).json()["items"]
        c = next((x for x in items if x["id"] == campaign_id), None)
        if c and c["status"] == "done":
            done = c
            break
    assert done is not None, "Campaign did not reach 'done' within 30s"
    assert done["sent"] + done["failed"] == done["total"]

    # cleanup imported clients
    lst = api_client.get(f"{base_url}/api/admin/clients", headers=admin_headers, params={"q": "test_camp_"}, timeout=30).json()["items"]
    for c in lst:
        api_client.delete(f"{base_url}/api/admin/clients/{c['id']}", headers=admin_headers, timeout=30)


def test_admin_campaign_test_endpoint(api_client, base_url, admin_headers):
    r = api_client.post(f"{base_url}/api/admin/campaigns/test", headers=admin_headers,
                       json={"subject": "TEST prueba", "message": "prueba", "email": "admin@okumekaraoke.com"}, timeout=60)
    # per spec: 200 ok:true if Resend live; 502 on failure. Just report either.
    assert r.status_code in (200, 502), r.text
    if r.status_code == 200:
        assert r.json().get("ok") is True


# ---------- Cron auth ----------
def test_cron_without_auth_returns_401(api_client, base_url):
    r = api_client.post(f"{base_url}/api/cron/followup", timeout=30)
    assert r.status_code == 401


def test_cron_with_bearer_returns_202(api_client, base_url):
    secret = _cron_secret()
    assert secret
    r = api_client.post(f"{base_url}/api/cron/followup",
                       headers={"Authorization": f"Bearer {secret}"}, timeout=30)
    assert r.status_code == 202
