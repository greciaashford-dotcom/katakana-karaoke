"""Catálogo de canciones de Karaoke Katakana: normalización, parseo y exportación Excel."""

import re
import unicodedata
import uuid
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CATALOG_VERSION = "katakana-2026-10-01"
SONG_NAMESPACE = uuid.UUID("6b1d5c0e-7a2f-4f37-9a0e-4b8f6b8d2a11")

# key, etiqueta, código del sistema original
LANGUAGES = [
    ("es", "Español", "ES"), ("en", "Inglés", "UK"), ("fr", "Francés", "FR"), ("pt", "Portugués", "PT"),
    ("it", "Italiano", "IT"), ("ca", "Catalán", "CA"), ("de", "Alemán", "DE"), ("gl", "Gallego", "GA"),
]
LANG_LABEL = {key: label for key, label, _ in LANGUAGES}
LANG_SOURCE = {key: code for key, _, code in LANGUAGES}
OTHER_LANG = ("xx", "Otros")


def _strip_accents(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value)
    return "".join(ch for ch in decomposed if not unicodedata.combining(ch))


def search_text(value: str) -> str:
    """Minúsculas, sin tildes ni signos: 'Don't Stop – AC/DC' -> 'dont stop ac dc'."""
    text = _strip_accents(str(value or "")).lower()
    text = re.sub(r"['´`’‘]", "", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return text.strip()


def search_tokens(query: str) -> list[str]:
    return [token for token in search_text(query).split(" ") if token][:8]


def collapse(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def resolve_lang(value) -> str:
    raw = search_text(value)
    if not raw:
        return ""
    for key, label, code in LANGUAGES:
        if raw in (key, search_text(label), code.lower()):
            return key
    aliases = {"ingles": "en", "english": "en", "uk": "en", "en": "en", "espanol": "es", "castellano": "es", "spanish": "es",
               "frances": "fr", "french": "fr", "portugues": "pt", "italiano": "it", "catalan": "ca", "aleman": "de", "gallego": "gl", "ga": "gl"}
    return aliases.get(raw, OTHER_LANG[0])


def lang_label(key: str) -> str:
    return LANG_LABEL.get(key, OTHER_LANG[1])


def build_search(artist: str, title: str, code: str) -> str:
    compact = re.sub(r"[^a-z0-9]", "", search_text(code))
    return collapse(f"{search_text(artist)} {search_text(title)} {search_text(code)} {compact}")


def song_document(row: dict, *, song_id: str | None = None) -> dict | None:
    """Construye el documento MongoDB de una canción a partir de una fila normalizada."""
    artist = collapse(row.get("artist"))[:200]
    title = collapse(row.get("title"))[:200]
    if not artist or not title:
        return None
    code = collapse(row.get("code"))[:60]
    try:
        source_id = int(str(row.get("sourceId") or "").strip())
    except ValueError:
        source_id = None
    lang = resolve_lang(row.get("lang")) or resolve_lang(row.get("language")) or "es"
    doc = {
        "id": song_id or (str(uuid.uuid5(SONG_NAMESPACE, f"katakana:{source_id}")) if source_id is not None else str(uuid.uuid4())),
        "sourceId": source_id, "artist": artist, "title": title, "code": code, "lang": lang, "language": lang_label(lang),
        "artistKey": search_text(artist), "titleKey": search_text(title), "codeKey": re.sub(r"[^a-z0-9]", "", search_text(code)),
        "search": build_search(artist, title, code),
    }
    return doc


# ---------- Excel ----------
HEADER_PATTERNS = [
    ("sourceId", r"^(id|n|num|numero|ref)$"),
    ("lang", r"codigo idioma|cod idioma|lang|language code"),
    ("code", r"codigo|code|cod$|referencia"),
    ("artist", r"artista|interprete|artist|cantante|grupo"),
    ("title", r"titulo|title|cancion|tema|song"),
    ("language", r"idioma|language|lengua"),
]


def detect_column(header) -> str | None:
    h = search_text(header)
    if not h:
        return None
    for key, pattern in HEADER_PATTERNS:
        if re.search(pattern, h):
            return key
    return None


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_catalog_workbook(data: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheets = workbook.worksheets
    # Preferimos la hoja que contenga cabeceras de artista y título (p. ej. "Canciones").
    for sheet in sheets:
        rows_iter = sheet.iter_rows(values_only=True)
        for _ in range(15):
            header = next(rows_iter, None)
            if header is None:
                break
            mapping: dict[int, str] = {}
            for index, value in enumerate(header):
                key = detect_column(value)
                if key and key not in mapping.values():
                    mapping[index] = key
            if "artist" in mapping.values() and "title" in mapping.values():
                rows = []
                for row in rows_iter:
                    record = {key: _cell(row[index]) if index < len(row) else "" for index, key in mapping.items()}
                    if record.get("artist") and record.get("title"):
                        rows.append(record)
                return rows
    return []


EXPORT_COLUMNS = [("ID", 9), ("Artista", 38), ("Título", 46), ("Código", 18), ("Idioma", 13), ("Código idioma", 15)]


def build_catalog_workbook(songs: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Canciones"
    sheet.append([header for header, _ in EXPORT_COLUMNS])
    for index, (_, width) in enumerate(EXPORT_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="FFE8461E")
    for s in songs:
        sheet.append([s.get("sourceId") or "", s.get("artist", ""), s.get("title", ""), s.get("code", ""), s.get("language", ""), LANG_SOURCE.get(s.get("lang", ""), "")])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}1"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
