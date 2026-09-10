import re
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .common import MADRID, normalize, parse_iso

COLUMNS = [
    ("Correo", "email", 34), ("Nombre", "name", 26), ("Teléfono", "phone", 18), ("Origen", "source", 16), ("Fecha de alta", "createdAt", 20),
    ("Suscrito", "subscribed", 10), ("Último correo", "lastEmailAt", 20), ("Próximo seguimiento", "nextFollowupAt", 20), ("Seguimientos", "followupCount", 13), ("Notas", "notes", 30),
]


def format_date(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return parse_iso(iso).astimezone(MADRID).strftime("%d/%m/%y, %H:%M")
    except ValueError:
        return str(iso)


def build_clients_workbook(clients: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clientes"
    sheet.append([header for header, _, _ in COLUMNS])
    for index, (_, _, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFF8EFD2")
        cell.fill = PatternFill(fill_type="solid", fgColor="FF29493A")
    sheet.row_dimensions[1].height = 22
    for c in clients:
        subscribed = c.get("subscribed") is not False
        sheet.append([
            c.get("email", ""), c.get("name") or "", c.get("phone") or "", c.get("source") or "", format_date(c.get("createdAt")),
            "Sí" if subscribed else "No", format_date(c.get("lastEmailAt")), format_date(c.get("nextFollowupAt")) if subscribed else "",
            c.get("followupCount") or 0, c.get("notes") or "",
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def detect_key(header) -> str | None:
    h = normalize(str(header or "")).strip()
    if not h:
        return None
    if re.search(r"correo|e-?mail|mail", h):
        return "email"
    if re.search(r"nombre|name", h):
        return "name"
    if re.search(r"tel|phone|movil|celular|whatsapp", h):
        return "phone"
    if re.search(r"origen|source", h):
        return "source"
    if re.search(r"suscri|subscri|activo", h):
        return "subscribed"
    if re.search(r"nota|comentario|notes", h):
        return "notes"
    return None


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def parse_clients_workbook(data: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.worksheets[0] if workbook.worksheets else None
    if sheet is None:
        return []
    rows_iter = sheet.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        return []
    mapping: dict[int, str] = {}
    for index, value in enumerate(header):
        key = detect_key(value)
        if key and key not in mapping.values():
            mapping[index] = key
    has_header = "email" in mapping.values()
    if not has_header:
        mapping = {0: "email"}
    rows = []
    source_rows = rows_iter if has_header else [header, *rows_iter]
    for row in source_rows:
        record = {key: cell_text(row[index]) if index < len(row) else "" for index, key in mapping.items()}
        if not record.get("email"):
            continue
        if "subscribed" in record:
            record["subscribed"] = not re.fullmatch(r"(no|false|0|baja)", record["subscribed"] or "", flags=re.IGNORECASE)
        rows.append(record)
    return rows
